"""
Genera el archivo Excel a partir de los movimientos extraídos de cualquier banco.
"""

from typing import Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from bancos.base import Movimiento, InfoCuenta

COLOR_DEPOSITO = "E8F5E9"
COLOR_RETIRO   = "FFEBEE"
COLOR_SALDO    = "E3F2FD"
COLOR_ALT_ROW  = "FAFAFA"
COLOR_WHITE    = "FFFFFF"
COLOR_TOTAL_D  = "C8E6C9"
COLOR_TOTAL_R  = "FFCDD2"

MONEY_FMT = '#,##0.00'

THIN = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)


def _cell(ws, row, col, value=None, *, bold=False, color=None, bg=None,
          align="left", fmt=None, wrap=False, size=9, fg="212121"):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(bold=bold, size=size, color=fg)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    c.border = THIN
    if bg:
        c.fill = PatternFill("solid", fgColor=bg)
    if fmt:
        c.number_format = fmt
    return c


# Columnas de la hoja: (encabezado, ancho, tipo). El tipo dice de dónde sale el
# valor y cómo se pinta. Col A = ID ESTABLE (llave de sync con la BD).
_COLS_GENERICO = [
    ("ID", 8, "id"), ("FECHA", 14, "fecha"), ("FOLIO", 12, "folio"),
    ("DESCRIPCION", 52, "desc"), ("DEPOSITO", 16, "dep"),
    ("RETIRO", 16, "ret"), ("SALDO", 18, "saldo"),
]
# BBVA: las mismas columnas del estado de cuenta (doble fecha y doble saldo).
_COLS_BBVA = [
    ("ID", 8, "id"), ("FECHA OPER", 11, "fecha"), ("FECHA LIQ", 11, "fecha_liq"),
    ("COD.", 7, "folio"), ("DESCRIPCIÓN", 32, "desc1"), ("REFERENCIA", 54, "ref"),
    ("CARGOS", 15, "ret"), ("ABONOS", 15, "dep"),
    ("SALDO OPERACIÓN", 17, "saldo"), ("SALDO LIQUIDACIÓN", 17, "saldo_liq"),
]
_MONEY = {"dep", "ret", "saldo", "saldo_liq"}


def _valor(mov: Movimiento, tipo: str):
    desc = mov.descripcion or ""
    if tipo == "id":
        return getattr(mov, "id_mov", None) or None
    if tipo == "desc":
        return desc.replace(" | ", "\n")
    if tipo == "desc1":                       # 1ª línea del movimiento
        return desc.partition(" | ")[0]
    if tipo == "ref":                         # líneas siguientes (RFC, AUT, Ref.)
        return desc.partition(" | ")[2].replace(" | ", "\n") or None
    return getattr(mov, {"dep": "deposito", "ret": "retiro"}.get(tipo, tipo), None)


def export_to_excel(
    movimientos: list[Movimiento],
    output_path: str,
    info: InfoCuenta,
    empresa_nombre: str = "",
    usuario: str = "",
) -> str:
    """
    Genera el Excel de movimientos.

    Parametros adicionales:
      empresa_nombre : nombre de la empresa obtenido del catalogo por RFC
      usuario        : nombre del usuario que realizo la conversion
    """
    color_header = info.banco and _get_bank_color(info.banco) or "2C3E50"

    cols   = _COLS_BBVA if (info.banco or "").upper() == "BBVA" else _COLS_GENERICO
    col_de = {tipo: i for i, (_, _, tipo) in enumerate(cols, 1)}
    ult    = get_column_letter(len(cols))
    # Las etiquetas de saldo inicial/totales ocupan de B hasta antes del 1er importe.
    n_etq   = min(col_de[t] for t in _MONEY if t in col_de) - 1
    etq_fin = get_column_letter(n_etq)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Movimientos"

    # ── Fila 1: Empresa (si se encontro) o Titulo del banco ───────────────────
    ws.merge_cells(f"A1:{ult}1")
    c1 = ws["A1"]
    if empresa_nombre:
        c1.value = empresa_nombre.upper()
    else:
        c1.value = f"Estado de Cuenta {info.banco}"
    c1.font = Font(bold=True, size=14, color="FFFFFF")
    c1.fill = PatternFill("solid", fgColor=color_header)
    c1.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # ── Fila 2: Banco + Periodo ────────────────────────────────────────────────
    ws.merge_cells(f"A2:{ult}2")
    c2 = ws["A2"]
    banco_periodo = f"{info.banco}  —  {info.periodo or ''}"
    if empresa_nombre and info.rfc:
        banco_periodo = f"RFC: {info.rfc}   |   {banco_periodo}"
    c2.value = banco_periodo
    c2.font = Font(bold=True, size=10, color="FFFFFF")
    c2.fill = PatternFill("solid", fgColor=color_header)
    c2.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    # ── Fila 3: Info de cuenta + usuario ──────────────────────────────────────
    ws.merge_cells(f"A3:{ult}3")
    c3 = ws["A3"]
    parts = []
    if info.numero_cuenta:
        parts.append(f"Cuenta: {info.numero_cuenta}")
    if info.titular:
        parts.append(f"Titular: {info.titular}")
    if not empresa_nombre and info.rfc:
        parts.append(f"RFC: {info.rfc}")
    parts.append(f"{len(movimientos)} movimientos")
    if usuario:
        parts.append(f"Generado por: {usuario}")
    c3.value = "   |   ".join(parts)
    c3.font = Font(italic=True, size=8, color="555555")
    c3.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[3].height = 14

    row_offset = 1   # filas extra por fila 2 adicional

    # ── Encabezados ───────────────────────────────────────────────────────────
    HDR_ROW = 4
    for col, (h, _, _) in enumerate(cols, 1):
        c = ws.cell(row=HDR_ROW, column=col, value=h)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=color_header)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = THIN
    ws.row_dimensions[HDR_ROW].height = 22

    # ── Saldo inicial ─────────────────────────────────────────────────────────
    data_start = HDR_ROW + 1
    if info.saldo_inicial is not None:
        ws.merge_cells(f"B{data_start}:{etq_fin}{data_start}")
        c = ws.cell(row=data_start, column=2, value="SALDO INICIAL DEL PERIODO ANTERIOR")
        c.font = Font(bold=True, italic=True, size=9, color="444444")
        c.alignment = Alignment(horizontal="right", vertical="center")
        c.border = THIN
        for col in range(3, len(cols) + 1):
            ws.cell(row=data_start, column=col).border = THIN
        _cell(ws, data_start, col_de["saldo"], info.saldo_inicial,
              fmt=MONEY_FMT, bg=COLOR_SALDO, bold=True, align="right", size=9)
        data_start += 1

    # ── Filas de datos ────────────────────────────────────────────────────────
    for i, mov in enumerate(movimientos):
        row = data_start + i
        alt = COLOR_ALT_ROW if i % 2 else COLOR_WHITE
        num_lineas = 1
        for col, (_, _, tipo) in enumerate(cols, 1):
            v = _valor(mov, tipo)
            if tipo in ("dep", "ret"):
                bg = (COLOR_DEPOSITO if tipo == "dep" else COLOR_RETIRO) if v else alt
                _cell(ws, row, col, v, fmt=MONEY_FMT, bg=bg, align="right")
            elif tipo in _MONEY:
                _cell(ws, row, col, v, fmt=MONEY_FMT, bg=COLOR_SALDO, align="right")
            elif tipo in ("desc", "desc1", "ref"):
                _cell(ws, row, col, v, bg=alt, wrap=True)
                num_lineas = max(num_lineas, str(v or "").count("\n") + 1)
            else:
                _cell(ws, row, col, v, bg=alt, align="center")
        ws.row_dimensions[row].height = max(18, num_lineas * 14)

    # ── Totales ────────────────────────────────────────────────────────────────
    total_row = data_start + len(movimientos)
    ws.merge_cells(f"B{total_row}:{etq_fin}{total_row}")
    c = ws.cell(row=total_row, column=2, value="TOTALES DEL PERIODO")
    c.font = Font(bold=True, size=10, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=color_header)
    c.alignment = Alignment(horizontal="right", vertical="center")
    c.border = THIN
    for col in [1] + list(range(3, n_etq + 1)):
        ws.cell(row=total_row, column=col).border = THIN
        ws.cell(row=total_row, column=col).fill = PatternFill("solid", fgColor=color_header)

    total_dep = info.total_depositos or sum(m.deposito or 0 for m in movimientos)
    total_ret = info.total_retiros   or sum(m.retiro   or 0 for m in movimientos)
    _cell(ws, total_row, col_de["dep"], total_dep, fmt=MONEY_FMT, bg=COLOR_TOTAL_D, bold=True, align="right")
    _cell(ws, total_row, col_de["ret"], total_ret, fmt=MONEY_FMT, bg=COLOR_TOTAL_R, bold=True, align="right")
    saldo_fin = info.saldo_final or (movimientos[-1].saldo if movimientos else None)
    _cell(ws, total_row, col_de["saldo"], saldo_fin, fmt=MONEY_FMT, bg=COLOR_SALDO, bold=True, align="right")
    if "saldo_liq" in col_de:
        liq_fin = next((m.saldo_liq for m in reversed(movimientos) if m.saldo_liq is not None), None)
        _cell(ws, total_row, col_de["saldo_liq"], liq_fin, fmt=MONEY_FMT, bg=COLOR_SALDO, bold=True, align="right")

    # ── Anchos ────────────────────────────────────────────────────────────────
    for col, (_, w, _) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.freeze_panes = f"A{data_start}"
    ws.auto_filter.ref = f"A{HDR_ROW}:{ult}{total_row - 1}"

    wb.save(output_path)
    return output_path


def _get_bank_color(banco: str) -> str:
    colors = {
        "Santander": "CC0000",
        "BBVA":      "004481",
        "Banregio":  "005A8E",
        "Banorte":   "006341",
        "Inbursa":   "E31837",
        "HSBC":      "DB0011",
        "Banamex":   "003DA5",
        "Scotiabank":"EC111A",
    }
    for key, color in colors.items():
        if key.lower() in banco.lower():
            return color
    return "2C3E50"
